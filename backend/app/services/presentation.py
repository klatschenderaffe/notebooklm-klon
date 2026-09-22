import io
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from app.services.design import DEFAULT_DESIGN

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
MARGIN = Inches(0.6)
ACCENT_BAR_HEIGHT = Inches(0.12)


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _set_background(slide: Any, hex_color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(hex_color)


def _add_accent_bar(slide: Any, hex_color: str) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, ACCENT_BAR_HEIGHT)
    bar.fill.solid()
    bar.fill.fore_color.rgb = _rgb(hex_color)
    bar.line.fill.background()
    bar.shadow.inherit = False


def _add_title_slide(prs: Any, title: str, design: dict[str, str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, design["background_color"])
    _add_accent_bar(slide, design["accent_color"])

    title_box = slide.shapes.add_textbox(MARGIN, Inches(3.0), SLIDE_WIDTH - 2 * MARGIN, Inches(1.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.text = title
    p = tf.paragraphs[0]
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.name = design["heading_font"]
    p.font.color.rgb = _rgb(design["text_color"])

    subtitle_box = slide.shapes.add_textbox(
        MARGIN, Inches(4.4), SLIDE_WIDTH - 2 * MARGIN, Inches(0.6)
    )
    stf = subtitle_box.text_frame
    stf.text = "Erstellt mit NotebookLM-Klon"
    sp = stf.paragraphs[0]
    sp.font.size = Pt(16)
    sp.font.name = design["body_font"]
    sp.font.color.rgb = _rgb(design["accent_color"])


def _add_content_slide(prs: Any, slide_data: dict[str, Any], design: dict[str, str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, design["background_color"])
    _add_accent_bar(slide, design["accent_color"])

    title_box = slide.shapes.add_textbox(MARGIN, Inches(0.4), SLIDE_WIDTH - 2 * MARGIN, Inches(0.9))
    ttf = title_box.text_frame
    ttf.word_wrap = True
    ttf.text = slide_data.get("title", "")
    tp = ttf.paragraphs[0]
    tp.font.size = Pt(28)
    tp.font.bold = True
    tp.font.name = design["heading_font"]
    tp.font.color.rgb = _rgb(design["text_color"])

    body_box = slide.shapes.add_textbox(MARGIN, Inches(1.6), SLIDE_WIDTH - 2 * MARGIN, Inches(5.3))
    body_tf = body_box.text_frame
    body_tf.word_wrap = True
    bullets = slide_data.get("bullets", [])
    for i, bullet in enumerate(bullets):
        p = body_tf.paragraphs[0] if i == 0 else body_tf.add_paragraph()
        p.text = f"•  {bullet}"
        p.font.size = Pt(18)
        p.font.name = design["body_font"]
        p.font.color.rgb = _rgb(design["text_color"])
        p.space_after = Pt(12)


def build_presentation(
    title: str,
    slides: list[dict[str, Any]],
    design: dict[str, str] | None = None,
) -> bytes:
    design = design or DEFAULT_DESIGN

    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    _add_title_slide(prs, title, design)

    if not slides:
        _add_content_slide(prs, {"title": "Keine Inhalte verfügbar.", "bullets": []}, design)
    else:
        for slide_data in slides:
            _add_content_slide(prs, slide_data, design)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
