import io

from pptx import Presentation

from app.services.presentation import build_presentation


def test_build_presentation_creates_title_and_content_slides() -> None:
    slides = [
        {"title": "Einleitung", "bullets": ["Punkt A", "Punkt B"]},
        {"title": "Fazit", "bullets": ["Zusammenfassung"]},
    ]

    pptx_bytes = build_presentation("Meine Präsentation", slides)
    prs = Presentation(io.BytesIO(pptx_bytes))

    assert len(prs.slides) == 3
    assert prs.slides[0].shapes.title.text == "Meine Präsentation"
    assert prs.slides[1].shapes.title.text == "Einleitung"
    assert prs.slides[2].shapes.title.text == "Fazit"


def test_build_presentation_without_slides_adds_placeholder() -> None:
    pptx_bytes = build_presentation("Leer", [])
    prs = Presentation(io.BytesIO(pptx_bytes))

    assert len(prs.slides) == 2
